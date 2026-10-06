"""
document_extractor.py - Text extraction from PDF and DOCX files.

Extracts text with pdfplumber (PDF) or python-docx (DOCX).
Falls back to OCR if a PDF is scanned (no selectable text).
"""

import io
import logging

logger = logging.getLogger(__name__)

# Magic bytes for file type validation
_PDF_MAGIC = b"%PDF"
_DOCX_MAGIC = b"PK\x03\x04"  # ZIP format (DOCX is a ZIP)

MAX_SIZE_BYTES = 5 * 1024 * 1024  # 5 MB


def validate_file(filename: str, content_type: str, data: bytes) -> str:
    """
    Validate file by extension, MIME type, and magic bytes.
    Returns the detected type: 'pdf' or 'docx'.
    Raises ValueError on invalid input.
    """
    if len(data) > MAX_SIZE_BYTES:
        raise ValueError(f"File too large ({len(data) // 1024} KB). Maximum allowed is 5 MB.")

    name_lower = filename.lower()
    mime_lower = content_type.lower()

    if name_lower.endswith(".pdf") or "pdf" in mime_lower:
        if not data.startswith(_PDF_MAGIC):
            raise ValueError("File claims to be a PDF but magic bytes don't match.")
        return "pdf"
    elif name_lower.endswith(".docx") or "word" in mime_lower or "openxmlformats" in mime_lower:
        if not data.startswith(_DOCX_MAGIC):
            raise ValueError("File claims to be a DOCX but magic bytes don't match.")
        return "docx"
    else:
        raise ValueError(
            f"Unsupported file type '{content_type}'. Only PDF and DOCX are accepted."
        )


def extract_text_from_pdf(data: bytes) -> tuple[str, dict]:
    """
    Extract text and metadata from PDF bytes using pdfplumber.
    Falls back to OCR if no selectable text is found (scanned PDF).
    Returns (text, metadata_dict).
    """
    import pdfplumber

    metadata: dict = {}
    text_parts: list[str] = []

    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            # Collect metadata
            if pdf.metadata:
                metadata = {k: str(v) for k, v in pdf.metadata.items() if v}

            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)

    except Exception as e:
        logger.warning("pdfplumber failed: %s", e)

    full_text = "\n".join(text_parts).strip()

    # Fallback: if no text extracted, treat as scanned — use OCR
    if not full_text:
        logger.info("No selectable text found in PDF. Attempting OCR fallback...")
        try:
            from backend.app.ocr import extract_text_from_image
            full_text = extract_text_from_image(data)
            metadata["ocr_fallback"] = "true"
        except Exception as e:
            logger.warning("OCR fallback failed: %s", e)
            full_text = ""

    return full_text, metadata


def extract_text_from_docx(data: bytes) -> tuple[str, dict]:
    """
    Extract text and basic metadata from DOCX bytes using python-docx.
    Returns (text, metadata_dict).
    """
    import docx

    try:
        doc = docx.Document(io.BytesIO(data))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

        # Also extract from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())

        # Core properties metadata
        props = doc.core_properties
        metadata = {
            "author": props.author or "",
            "created": str(props.created) if props.created else "",
            "modified": str(props.modified) if props.modified else "",
            "last_modified_by": props.last_modified_by or "",
        }

        return "\n".join(paragraphs), metadata

    except Exception as e:
        logger.warning("python-docx extraction failed: %s", e)
        return "", {}
